#!/usr/bin/env python3
"""校舎配置図（Excel）から js/data.js を作るスクリプト。

使い方:
    python3 tools/excel_to_data.py 校舎配置図.xlsx js/data.js

Excel の1枚目のシートに 1階（左上）・2階（左下）・3階（右下）の図面が
セルの結合で描かれている前提です。部屋名が入った結合セルの範囲を
そのまま部屋の位置・大きさとして読み取ります。
"""
import json
import re
import sys
import unicodedata

import openpyxl

# 2階・3階の図面を1階の図面と同じ位置に重ねるためのずらし量（セル数）
FLOOR2_ROW_OFFSET = 72
FLOOR3_COL_OFFSET = 83

# 図面の見出しなど、部屋ではない文字
SKIP = {"１階", "２階", "３階", "神奈川県立麻生高等学校", "令和８年度　校舎配置図", "管理運営グループ"}


def floor_of(row, col):
    """セルの位置から (階, 1階基準の行, 1階基準の列) を返す。"""
    if row <= 70 and col <= 75:
        return 1, row, col
    if row >= 72 and col <= 75:
        return 2, row - FLOOR2_ROW_OFFSET, col
    if row >= 72 and col >= 85:
        return 3, row - FLOOR2_ROW_OFFSET, col - FLOOR3_COL_OFFSET
    return None


def clean(text):
    text = text.replace("\n", "")
    text = re.sub(r"[　 ]+", "", text)
    m = re.match(r"教室([０-９0-9])－([０-９0-9])", text)
    if m:
        g = unicodedata.normalize("NFKC", m.group(1))
        c = unicodedata.normalize("NFKC", m.group(2))
        return f"{g}年{c}組 教室", {"grade": int(g), "cls": int(c)}
    if text == "ＳＴＵＤＹＨＡＬＬ":
        text = "STUDY HALL"
    if text == "光庭":
        text = "光庭（中庭）"
    return text, {}


def main(src, dst):
    wb = openpyxl.load_workbook(src)
    ws = wb.worksheets[0]
    merged = {(m.min_row, m.min_col): m for m in ws.merged_cells.ranges}
    rooms = []
    for row in ws.iter_rows():
        for cell in row:
            if not isinstance(cell.value, str) or not cell.value.strip():
                continue
            if cell.value.strip() in SKIP:
                continue
            m = merged.get((cell.row, cell.column))
            r0, c0, r1, c1 = (m.min_row, m.min_col, m.max_row, m.max_col) if m else (cell.row, cell.column, cell.row, cell.column)
            f = floor_of(r0, c0)
            if not f:
                continue
            floor, rr, cc = f
            name, extra = clean(cell.value)
            room = {
                "id": f"F{floor}-{cell.coordinate}",
                "name": name,
                "floor": floor,
                "c0": cc, "r0": rr,
                "c1": cc + (c1 - c0), "r1": rr + (r1 - r0),
            }
            room.update(extra)
            rooms.append(room)
    with open(dst, "w", encoding="utf-8") as fp:
        fp.write("// tools/excel_to_data.py で Excel の校舎配置図から自動生成した部屋データ\n")
        fp.write("// c0,c1 = 列番号（左→右）、r0,r1 = 行番号（上→下）。1階図面のセル位置に揃えてある\n")
        fp.write("export const ROOMS_RAW = ")
        json.dump(rooms, fp, ensure_ascii=False, indent=0)
        fp.write(";\n")
    print(f"{len(rooms)} rooms -> {dst}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
